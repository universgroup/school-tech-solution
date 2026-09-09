"""
Usage :
    python manage.py import_classes /chemin/vers/Classe.xlsx --dry-run
    python manage.py import_classes /chemin/vers/Classe.xlsx

Fichier attendu : colonnes IDclasse, nom_classe, IDCycle, tranche1, tranche2,
Frais_scolarité, frais_inscription, frais_reinscription

Place ce fichier dans <app_name>/management/commands/import_classes.py
TODO: remplacer <app_name> par le nom réel de l'app contenant Classe.

PRE-REQUIS : import_cycles doit avoir été exécuté avant (IDCycle référence
les id CycleScolaire créés avec les mêmes identifiants legacy).

tranche3 n'est pas présent dans le fichier : laissé à sa valeur par défaut (0).

CIBLER UNE BASE PRECISE : par défaut Django écrit sur l'alias 'default' de
settings.DATABASES. Précisez l'alias exact si besoin :
    python manage.py import_classes fichier.xlsx --database=production
"""

from decimal import Decimal, InvalidOperation

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

import openpyxl

from gAdministration.models import Classe, CycleScolaire 


def to_decimal(value):
    if value in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(value))
    except InvalidOperation:
        raise ValueError(f"Montant invalide: {value!r}")


class Command(BaseCommand):
    help = "Importe Classe.xlsx vers le modèle Classe."

    def add_arguments(self, parser):
        parser.add_argument("fichier", type=str)
        parser.add_argument("--dry-run", action="store_true")
        parser.add_argument(
            "--database",
            type=str,
            default="default",
            help="Alias de base cible tel que défini dans settings.DATABASES (défaut: 'default').",
        )

    def handle(self, *args, **options):
        db_alias = options["database"]

        wb = openpyxl.load_workbook(options["fichier"], data_only=True)
        ws = wb.active

        headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
        col = {name: idx for idx, name in enumerate(headers)}
        required_cols = [
            "IDclasse", "nom_classe", "IDCycle", "tranche1", "tranche2",
            "Frais_scolarité", "frais_inscription", "frais_reinscription",
        ]
        missing = [c for c in required_cols if c not in col]
        if missing:
            raise CommandError(f"Colonnes manquantes: {missing}")

        created, updated, errors = 0, 0, []
        with transaction.atomic(using=db_alias):
            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                def val(name):
                    return row[col[name]]

                legacy_id = val("IDclasse")
                nom_classe = str(val("nom_classe") or "").strip()
                idcycle_id = val("IDCycle")
                if legacy_id is None or not nom_classe or idcycle_id is None:
                    errors.append((row_num, "ligne incomplète — ignorée"))
                    continue

                if not CycleScolaire.objects.using(db_alias).filter(id=idcycle_id).exists():
                    errors.append((
                        row_num,
                        f"CycleScolaire id={idcycle_id} introuvable — lancez d'abord import_cycles",
                    ))
                    continue

                try:
                    defaults = {
                        "nom_classe": nom_classe,
                        "idcycle_id": idcycle_id,
                        "tranche1": to_decimal(val("tranche1")),
                        "tranche2": to_decimal(val("tranche2")),
                        "frais_scolarite": to_decimal(val("Frais_scolarité")),
                        "frais_inscription": to_decimal(val("frais_inscription")),
                        "frais_reinscription": to_decimal(val("frais_reinscription")),
                    }
                except ValueError as exc:
                    errors.append((row_num, str(exc)))
                    continue

                if options["dry_run"]:
                    self.stdout.write(f"[DRY-RUN] id={legacy_id} {defaults}")
                    continue

                obj, was_created = Classe.objects.using(db_alias).update_or_create(
                    id=legacy_id, defaults=defaults
                )
                created += int(was_created)
                updated += int(not was_created)

            if options["dry_run"]:
                transaction.set_rollback(True, using=db_alias)

        self.stdout.write(self.style.SUCCESS(f"[{db_alias}] Créés: {created} | Mis à jour: {updated}"))
        if errors:
            self.stdout.write(self.style.ERROR(f"{len(errors)} ligne(s) en erreur:"))
            for row_num, msg in errors:
                self.stdout.write(f"  ligne {row_num}: {msg}")