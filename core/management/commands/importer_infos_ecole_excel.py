"""
Usage :
    python manage.py import_ecole /chemin/vers/Ecole.xlsx --dry-run
    python manage.py import_ecole /chemin/vers/Ecole.xlsx

Fichier attendu : colonnes IDecole, Nom_ecole, ville_ecole, agrement_ecole,
pref_ou_commune, DSEE, BP_ecole, telephone1, telephone2, email_ecole,
site_internet, devise, logo, dg, comptable

Place ce fichier dans <app_name>/management/commands/import_ecole.py
TODO: remplacer <app_name> par le nom réel de l'app contenant Ecole.

NON IMPORTÉ (absents du fichier, laissés vides — champs nullable) :
dga, coordo_primaire, coordo_secondaire, coordo_maternelle, signa_dg,
signa_de, delai_tranche1, delai_tranche2, delai_reinscription, logo_ecole
(FileField non gérable depuis une valeur texte).

Les valeurs "X" (placeholder vu dans devise/comptable) sont traitées comme
vides plutôt qu'importées littéralement.

CIBLER UNE BASE PRECISE : par défaut Django écrit sur l'alias 'default' de
settings.DATABASES. Précisez l'alias exact si besoin :
    python manage.py import_ecole fichier.xlsx --database=production
"""

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

import openpyxl

from gAdministration.models import Ecole 


def clean_text(value, default=""):
    if value is None:
        return default
    text = str(value).strip()
    if text.upper() == "X":
        return default
    return text


class Command(BaseCommand):
    help = "Importe Ecole.xlsx vers le modèle Ecole."

    def add_arguments(self, parser):
        parser.add_argument("fichier", type=str)
        parser.add_argument("--dry-run", action="store_true")
        parser.add_argument(
            "--database",
            type=str,
            default="default",
            help="Alias de base cible tel que défini dans settings.DATABASES (défaut: 'default').",
        )

    def handle(self, *args, **options):
        db_alias = options["database"]

        wb = openpyxl.load_workbook(options["fichier"], data_only=True)
        ws = wb.active

        headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
        col = {name: idx for idx, name in enumerate(headers)}
        required_cols = [
            "IDecole", "Nom_ecole", "ville_ecole", "agrement_ecole",
            "pref_ou_commune", "DSEE", "BP_ecole", "telephone1", "telephone2",
            "email_ecole", "site_internet", "devise", "logo", "dg", "comptable",
        ]
        missing = [c for c in required_cols if c not in col]
        if missing:
            raise CommandError(f"Colonnes manquantes: {missing}")

        created, updated = 0, 0
        with transaction.atomic(using=db_alias):
            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                def val(name):
                    return row[col[name]]

                legacy_id = val("IDecole")
                nom_ecole = clean_text(val("Nom_ecole"))
                if legacy_id is None or not nom_ecole:
                    self.stdout.write(self.style.WARNING(f"Ligne {row_num} ignorée (incomplète)"))
                    continue

                defaults = {
                    "nom_ecole": nom_ecole,
                    "ville_ecole": clean_text(val("ville_ecole")),
                    "agrement_ecole": clean_text(val("agrement_ecole")),
                    "prefect_commune": clean_text(val("pref_ou_commune")),
                    "dsee": clean_text(val("DSEE")),
                    "bp_ecole": clean_text(val("BP_ecole"), default=None) or None,
                    "telephone1": clean_text(val("telephone1")),
                    "telephone2": clean_text(val("telephone2"), default=None) or None,
                    "email_ecole": clean_text(val("email_ecole")),
                    "site_internet": clean_text(val("site_internet"), default=None) or None,
                    "devise_ecole": clean_text(val("devise")),
                    "dg": clean_text(val("dg")),
                    "comptable": clean_text(val("comptable"), default=None) or None,
                }

                if options["dry_run"]:
                    self.stdout.write(f"[DRY-RUN] id={legacy_id} {defaults}")
                    continue

                obj, was_created = Ecole.objects.using(db_alias).update_or_create(
                    id=legacy_id, defaults=defaults
                )
                created += int(was_created)
                updated += int(not was_created)

            if options["dry_run"]:
                transaction.set_rollback(True, using=db_alias)

        self.stdout.write(self.style.SUCCESS(f"[{db_alias}] Créés: {created} | Mis à jour: {updated}"))
