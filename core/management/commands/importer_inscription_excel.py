"""
Usage :
    python manage.py import_inscriptions /chemin/vers/Inscription.xlsx --dry-run
    python manage.py import_inscriptions /chemin/vers/Inscription.xlsx

Fichier attendu (version épurée) : colonnes ID_inscription, IDannee,
IDCycle, IDclasse, Matricule, date_inscription, etat_inscription.

IDannee, IDCycle et IDclasse sont désormais de vrais identifiants
numériques legacy — ils correspondent directement aux clés primaires
préservées par import_annee_scolaire, import_cycles et import_classes.
Plus besoin de table de correspondance (CLASSE_MAPPING) : on fait des
lookups directs par id.

Place ce fichier dans <app_name>/management/commands/import_inscriptions.py
TODO: remplacer <app_name> par le nom réel de l'app contenant Inscription.

PRE-REQUIS : import_annee_scolaire, import_cycles, import_classes et
l'import des Eleve doivent avoir été exécutés avant (avec les mêmes id
legacy).

--- date_inscription : particularité auto_now ------------------------------
Le modèle Inscription définit date_inscription = DateField(auto_now=True),
donc Django écrase toujours cette valeur avec la date du jour au moment du
.save(). Pour préserver la date historique du fichier, ce script crée
l'objet puis force la vraie date via un .update() en queryset (qui
contourne auto_now).

CIBLER UNE BASE PRECISE : par défaut Django écrit sur l'alias 'default' de
settings.DATABASES. Précisez l'alias exact si besoin :
    python manage.py import_inscriptions fichier.xlsx --database=production
"""

import datetime

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

import openpyxl

from gAdministration.models import AnneeScolaire, Classe, CycleScolaire
from gEleve.models import Eleve, Inscription  


def parse_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Format de date non reconnu: {value!r}")


class Command(BaseCommand):
    help = "Importe Inscription.xlsx (version épurée) vers le modèle Inscription."

    def add_arguments(self, parser):
        parser.add_argument("fichier", type=str)
        parser.add_argument("--dry-run", action="store_true")
        parser.add_argument(
            "--database",
            type=str,
            default="default",
            help="Alias de base cible tel que défini dans settings.DATABASES (défaut: 'default').",
        )

    def handle(self, *args, **options):
        db_alias = options["database"]
        dry_run = options["dry_run"]

        wb = openpyxl.load_workbook(options["fichier"], data_only=True)
        ws = wb.active

        headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
        col = {name: idx for idx, name in enumerate(headers)}
        required_cols = [
            "ID_inscription", "IDannee", "IDCycle", "IDclasse", "Matricule",
            "date_inscription", "etat_inscription",
        ]
        missing = [c for c in required_cols if c not in col]
        if missing:
            raise CommandError(f"Colonnes manquantes: {missing}")

        created, updated, errors = 0, 0, []

        with transaction.atomic(using=db_alias):
            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                def val(name):
                    return row[col[name]]

                legacy_id = val("ID_inscription")
                matricule = str(val("Matricule") or "").strip()
                idannee = val("IDannee")
                idcycle = val("IDCycle")
                idclasse = val("IDclasse")
                etat = str(val("etat_inscription") or "").strip()

                if not all([legacy_id, matricule, idannee, idcycle, idclasse, etat]):
                    errors.append((row_num, "ligne incomplète — ignorée"))
                    continue

                try:
                    eleve = Eleve.objects.using(db_alias).get(matricule=matricule)
                except Eleve.DoesNotExist:
                    errors.append((row_num, f"Eleve matricule={matricule!r} introuvable"))
                    continue

                try:
                    classe = Classe.objects.using(db_alias).get(id=idclasse)
                except Classe.DoesNotExist:
                    errors.append((row_num, f"Classe id={idclasse!r} introuvable — lancez import_classes d'abord"))
                    continue

                try:
                    cycle = CycleScolaire.objects.using(db_alias).get(id=idcycle)
                except CycleScolaire.DoesNotExist:
                    errors.append((row_num, f"CycleScolaire id={idcycle!r} introuvable — lancez import_cycles d'abord"))
                    continue

                try:
                    annee = AnneeScolaire.objects.using(db_alias).get(id=idannee)
                except AnneeScolaire.DoesNotExist:
                    errors.append((row_num, f"AnneeScolaire id={idannee!r} introuvable — lancez import_annee_scolaire d'abord"))
                    continue

                try:
                    date_inscription = parse_date(val("date_inscription"))
                except ValueError as exc:
                    errors.append((row_num, str(exc)))
                    continue

                if dry_run:
                    self.stdout.write(
                        f"[DRY-RUN] id={legacy_id} eleve={matricule} classe_id={idclasse} "
                        f"cycle_id={idcycle} annee_id={idannee} etat={etat} date={date_inscription}"
                    )
                    continue

                obj, was_created = Inscription.objects.using(db_alias).update_or_create(
                    id=legacy_id,
                    defaults={
                        "mateleve": eleve,
                        "idclasse": classe,
                        "idcycle": cycle,
                        "annee_scolaire": annee,
                        "etat_inscription": etat,
                    },
                )
                # force la date historique (contourne auto_now du modèle)
                if date_inscription is not None:
                    Inscription.objects.using(db_alias).filter(pk=obj.pk).update(
                        date_inscription=date_inscription
                    )

                created += int(was_created)
                updated += int(not was_created)

            if dry_run:
                transaction.set_rollback(True, using=db_alias)

        self.stdout.write(self.style.SUCCESS(f"[{db_alias}] Créés: {created} | Mis à jour: {updated}"))
        if errors:
            self.stdout.write(self.style.ERROR(f"{len(errors)} ligne(s) en erreur:"))
            for row_num, msg in errors:
                self.stdout.write(f"  ligne {row_num}: {msg}")