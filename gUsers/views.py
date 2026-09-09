# Create your views here.
from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required
from django.contrib.auth.views import LoginView, LogoutView, PasswordChangeView, PasswordChangeDoneView
from django.shortcuts import render, redirect
from django.contrib import messages
from .forms import ConnexionForm, CreationUtilisateurForm, MotDePasseOublieForm, PhotoProfilForm, PasswordChangeCustomForm
from django.core.paginator import Paginator
from .models import NIVEAU_ACCES_CHOICES
from django.contrib.auth import views as auth_views
from django.urls import reverse_lazy
from django.db.models import Sum, Count, Q
from django.db.models.functions import TruncMonth
from django.contrib.auth.mixins import LoginRequiredMixin # Pour les vues basées sur une classe, on utilise LoginRequiredMixin (pas le décorateur @login_required)

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet


from gComptabilite.models import EtatPaiementTranche, Caisse, TYPE_OPERATION_CAISSE_CHOICES
from gAdministration.models import Ecole, Classe
from gEleve.models import Inscription, ETAT_INSCRIPTION, SEXE_ELEVE_CHOICES
from gComptabilite.views import affichersoldecaisse
from gPersonnel.models import Personnel
from core.context_processor import annee_scolaire_actuelle
from core.services import get_progression_paiements, get_totaux_paiements

from gUsers.decorators import action_requise

Utilisateur = get_user_model()


MOIS_FR = {1:'Jan',2:'Fév',3:'Mar',4:'Avr',5:'Mai',6:'Juin',7:'Juil',8:'Août',9:'Sep',10:'Oct',11:'Nov',12:'Déc'}


class ConnexionView(LoginView):
    authentication_form = ConnexionForm
    redirect_authenticated_user = True


class DeconnexionView(LogoutView):
    next_page = 'connexion'


# def _est_directeur_general(user):
#     return user.is_authenticated and user.est_directeur_general()


@action_requise('menu_administration')
def listeutilisateurs(request):
    use = Utilisateur.objects.all().order_by('id')

    pagineuser = Paginator(use, 10)
    numpageuser = request.GET.get('page')
    use = pagineuser.get_page(numpageuser)

    return render(request, 'gUsers/liste_utilisateurs.html', {'utilisateurs': use})


@action_requise('menu_administration')
def creerutilisateur(request):
    if request.method == 'POST':
        form = CreationUtilisateurForm(request.POST, request.FILES)
        if form.is_valid():
            form.save()
            messages.success(request, "Compte utilisateur créé avec succès.")
            form = CreationUtilisateurForm()
        else:
            messages.error(request, 'Vos mots de passe ne sont pas conformes. Veuillez reprendre à nouveau.')
    else:
        form = CreationUtilisateurForm()
    return render(request, 'gUsers/creer_utilisateur.html', {'form': form})

@action_requise('menu_administration')
def editerutilisateur(request, iduser):
    use = Utilisateur.objects.get(id=iduser)
    return render(request,'gUsers/modifier_utilisateur.html',dict(user=use, niveau_acces=NIVEAU_ACCES_CHOICES))



@login_required
def affichermonprofil(request, iduser):
    use = Utilisateur.objects.get(id=iduser)   
    return render(request, 'gUsers/mon_profil.html',dict(user=use))



@action_requise('menu_administration')
def modifierutilisateur(request, pk):
    if request.method == 'POST':
            use = Utilisateur.objects.get(id=pk)
            use.username = request.POST['username']
            use.first_name = request.POST['first_name']
            use.last_name = request.POST['last_name']
            use.email = request.POST['email_user']
            use.niveau_acces = request.POST['niveau_acces']

            if request.FILES.get('new_photo_profil'):
                use.photo_profil = request.FILES.get('new_photo_profil')

            if request.POST.get('passwrd'):
                use.password = request.POST['passwrd']
            use.save()
            return redirect('liste_utilisateurs')
    else:
        return redirect('../liste_utilisateurs/')


@action_requise('menu_administration')
def supprimerutilisateur(request, iduser):
    use = Utilisateur.objects.get(id=iduser)
    use.delete()
    return redirect('../liste_utilisateurs/')


eff_total_inscrits = 0
eff_total_reinscrits = 0
eff_garcon = 0
eff_fille = 0
eff_personnel = 0
t_entree = 0
t_sortie = 0
solde = 0
montant_restant = 0
annee_encours = {}

@login_required
def home(request):

    global eff_total_inscrits
    global eff_total_reinscrits
    global eff_garcon
    global eff_fille
    global eff_personnel
    global t_entree
    global t_sortie
    global solde
    global montant_restant
    global annee_encours

    liste_eleves = {}
    liste_personnel = {}
    liste_operation = {}
    totaux_paiements = {}

    liste_eleves = Inscription.objects.none()
    liste_personnel = Personnel.objects.none()
    liste_operation = Caisse.objects.none()

    annee_encours = annee_scolaire_actuelle(request)['annee_scolaire'] # annee_scolaire_actuelle vient du core/context_processor/

    context = {
            'progression_paiements': [],
            'inscriptions_labels': [],
            'inscriptions_data': [],
            'classes_labels': [],
            'classes_data': [],
            'total_inscrits': [],
            'total_garcons':[],
            'total_filles':[],
            'total_personnel':[],
            'total_reinscrits':[],
            'total_entree':[],
            'total_sortie':[],
            'solde_disponible':[],

        }

    # 1. Calcul de l'effectif total des inscrits et reinscrits, des garçons et filles, du personnel, des recettes et depenses

    liste_eleves = Inscription.objects.select_related('annee_scolaire','mateleve').filter(annee_scolaire__descript_annee__exact=annee_encours).all()

    eff_total_inscrits = liste_eleves.filter(etat_inscription__exact=ETAT_INSCRIPTION[0][1]).count()

    eff_total_reinscrits = liste_eleves.filter(etat_inscription__exact=ETAT_INSCRIPTION[1][0]).count()

    eff_garcon = liste_eleves.filter(Q(etat_inscription__exact=ETAT_INSCRIPTION[0][1]) | Q(etat_inscription__exact=ETAT_INSCRIPTION[1][0]),Q(mateleve__sexe_eleve__exact=SEXE_ELEVE_CHOICES[1][0])).count()

    eff_fille = liste_eleves.filter(Q(etat_inscription__exact=ETAT_INSCRIPTION[0][1]) | Q(etat_inscription__exact=ETAT_INSCRIPTION[1][0]),Q(mateleve__sexe_eleve__exact=SEXE_ELEVE_CHOICES[2][0])).count()

    liste_personnel = Personnel.objects.all()

    eff_personnel = liste_personnel.count()

    liste_operation = Caisse.objects.filter(anscolaire__descript_annee__exact=annee_encours).all()

    total_e = liste_operation.filter(type_operation__exact=TYPE_OPERATION_CAISSE_CHOICES[1][0]).aggregate(total_recette=Sum('montant_encaisse')) # Cumul des recettes

    total_s = liste_operation.filter(type_operation__exact=TYPE_OPERATION_CAISSE_CHOICES[2][0]).aggregate(total_depense=Sum('montant_encaisse')) # Cumul des dépenses

    t_entree = total_e['total_recette'] if total_e['total_recette'] is not None else 0
    t_sortie = total_s['total_depense'] if total_s['total_depense'] is not None else 0

    solde = affichersoldecaisse()
    

    # ---------- 1. Tableau "Progression des paiements de scolarité" ----------
    progression_paiements = get_progression_paiements(annee_encours)
    totaux_paiements = get_totaux_paiements(progression_paiements)

    # ---------- 2. Graphique "Évolution des inscriptions" ----------
    inscriptions_mois = (
        Inscription.objects
        .filter(annee_scolaire__descript_annee__exact=annee_encours)
        .annotate(mois=TruncMonth('date_inscription'))
        .values('mois')
        .annotate(total=Count('id'))
        .order_by('mois')
    )

    labels_inscriptions = []
    cumul_inscriptions = []
    total = 0
    for entree in inscriptions_mois:
        total += entree['total']
        labels_inscriptions.append(MOIS_FR[entree['mois'].month])
        cumul_inscriptions.append(total)

    # ---------- 3. Graphique "Répartition par classe" ----------
    repartition_classes = (
        Inscription.objects
        .filter(annee_scolaire__descript_annee__exact=annee_encours)
        .values('idclasse__nom_classe', 'idclasse__idcycle__cycle')
        .annotate(total=Count('id'))
        .order_by('idclasse_id')
    )

    context = {
        'progression_paiements': progression_paiements,
        'totaux_paiements': totaux_paiements,
        'inscriptions_labels': labels_inscriptions,         
        'inscriptions_data': cumul_inscriptions,
        'classes_labels': [c['idclasse__nom_classe'] for c in repartition_classes],
        'classes_cycles': [c['idclasse__idcycle__cycle'] for c in repartition_classes],
        'classes_data': [c['total'] for c in repartition_classes],
        'total_inscrits': eff_total_inscrits,
        'total_garcons': eff_garcon,
        'total_filles': eff_fille,
        'total_personnel': eff_personnel,
        'total_reinscrits': eff_total_reinscrits,
        'total_entree': t_entree,
        'total_sortie': t_sortie,
        'solde_disponible': solde,
    }

    return render(request,'global/dashboard.html', context)


# Permet de personnaliser le mail envoyé à l'utilisateur pour la reinitialisation de son mot de passe en y ajoutant le nom de l'école comme signature
class MotDePasseOublieView(auth_views.PasswordResetView):
    template_name = "gUsers/mot_de_passe_oublie.html"
    email_template_name = "emails/reset_mot_de_passe.txt"
    html_email_template_name = "emails/reset_mot_de_passe.html"
    subject_template_name = "emails/reset_mot_de_passe_sujet.txt"
    success_url = reverse_lazy('password_reset_done')
    form_class = MotDePasseOublieForm

    def form_valid(self, form):
        # Recalculé à chaque requête, jamais au démarrage du serveur.
        ecole = Ecole.objects.first()
        self.extra_email_context = {
            "nom_ecole": ecole.nom_ecole if ecole else "Ecole les Champions-Service Scolarité"
        }
        return super().form_valid(form)


@login_required
def changer_photo_profil(request):
    if request.method == 'POST':
        form = PhotoProfilForm(request.POST, request.FILES, instance=request.user)
        if form.is_valid():
            form.save()
            messages.success(request, "Photo de profil mise à jour avec succès.")
        else:
            messages.error(request, "Le fichier envoyé n'est pas valide.")

    return redirect(request.META.get('HTTP_REFERER', 'accueil'))



class PasswordChangeCustomView(LoginRequiredMixin, PasswordChangeView):
    form_class = PasswordChangeCustomForm
    template_name = 'gUsers/password_change_form.html'
    success_url = reverse_lazy('password_change_done')

    def form_valid(self, form):
        messages.success(self.request, 'Votre mot de passe a été modifié avec succès !')
        return super().form_valid(form)


class PasswordChangeDoneCustomView(LoginRequiredMixin, PasswordChangeDoneView):
    template_name = 'gUsers/password_change_done.html'


@login_required
def export_paiements_excel(request):

    annee_encours = annee_scolaire_actuelle(request)['annee_scolaire']
    progression_paiements = get_progression_paiements(annee_encours) if annee_encours is not None else []
    totaux = get_totaux_paiements(progression_paiements) if progression_paiements else None

    wb = Workbook()
    ws = wb.active
    ws.title = "Progression des paiements"

    entetes = ["Classe", "Montant attendu", "1ère tranche", "2ème tranche", "Montant restant"]
    ws.append(entetes)

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for ligne in progression_paiements:
        ws.append([
            ligne['nom_classe'],
            float(ligne['montant_attendu']),
            float(ligne['premiere_tranche']),
            float(ligne['deuxieme_tranche']),
            float(ligne['montant_restant']),
        ])

    if totaux is not None:
        ws.append([
            "TOTAUX",
            float(totaux['montant_attendu']),
            float(totaux['premiere_tranche']),
            float(totaux['deuxieme_tranche']),
            float(totaux['montant_restant']),
        ])
        ligne_total = ws.max_row
        for cell in ws[ligne_total]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")

    for colonne in ws.columns:
        largeur = max(len(str(cell.value)) for cell in colonne if cell.value is not None) + 4
        ws.column_dimensions[colonne[0].column_letter].width = largeur

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response['Content-Disposition'] = 'attachment; filename="tableau_progression_paiements.xlsx"'
    wb.save(response)
    return response


@login_required
def export_paiements_pdf(request):

    annee_encours = annee_scolaire_actuelle(request)['annee_scolaire']
    progression_paiements = get_progression_paiements(annee_encours) if annee_encours is not None else []
    totaux = get_totaux_paiements(progression_paiements) if progression_paiements else None

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="tableau_progression_paiements.pdf"'

    doc = SimpleDocTemplate(response, pagesize=landscape(A4), topMargin=1.5*cm, bottomMargin=1.5*cm)
    styles = getSampleStyleSheet()
    story = [Paragraph("Progression des paiements de scolarité", styles['Title']), Spacer(1, 12)]

    data = [["Classe", "Montant attendu", "1ère tranche", "2ème tranche", "Montant restant"]]
    for ligne in progression_paiements:
        data.append([
            ligne['nom_classe'],
            f"{ligne['montant_attendu']:,.0f}",
            f"{ligne['premiere_tranche']:,.0f}",
            f"{ligne['deuxieme_tranche']:,.0f}",
            f"{ligne['montant_restant']:,.0f}",
        ])

    ligne_total_index = None
    if totaux is not None:
        data.append([
            "TOTAUX",
            f"{totaux['montant_attendu']:,.0f}",
            f"{totaux['premiere_tranche']:,.0f}",
            f"{totaux['deuxieme_tranche']:,.0f}",
            f"{totaux['montant_restant']:,.0f}",
        ])
        ligne_total_index = len(data) - 1  # dernière ligne du tableau

    table = Table(data, repeatRows=1)
    style_commands = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4f46e5')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#dddddd')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f7f7fb')]),
        ('ALIGN', (1, 0), (-1, -1), 'RIGHT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]

    if ligne_total_index is not None:
        style_commands += [
            ('BACKGROUND', (0, ligne_total_index), (-1, ligne_total_index), colors.HexColor('#0f172a')),
            ('TEXTCOLOR', (0, ligne_total_index), (-1, ligne_total_index), colors.white),
            ('FONTNAME', (0, ligne_total_index), (-1, ligne_total_index), 'Helvetica-Bold'),
        ]

    table.setStyle(TableStyle(style_commands))
    story.append(table)
    doc.build(story)
    return response